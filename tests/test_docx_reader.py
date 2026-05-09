from pathlib import Path

from docx import Document

from src.readers.docx_reader import parse_docx_text, read_doc_text, read_docx_text
from src.normalizers import normalize_date, normalize_phone, normalize_text


def test_normalize_phone_removes_non_digits():
    assert normalize_phone("147 1469 1425") == "14714691425"


def test_normalize_text_collapses_whitespace():
    assert normalize_text("  乳源\n东阳光  ") == "乳源 东阳光"


def test_normalize_date_supports_spaced_chinese_digits():
    assert normalize_date("20 26 年 04 月 10 日").isoformat() == "2026-04-10"


def test_parse_docx_extracts_owner_fields_and_acceptance_range():
    text = (
        "项目编号：BHE-25030367/01 "
        "项目全称 乳源东阳光项目 "
        "用户姓名 黄汉民 "
        "联系电话14714691425 "
        "竣工验收2026年04月09日2026年04月10日"
    )

    parsed = parse_docx_text(text)

    assert parsed.project_code == "BHE-25030367/01"
    assert parsed.project_name == "乳源东阳光项目"
    assert parsed.contact_name == "黄汉民"
    assert parsed.contact_phone == "14714691425"
    assert parsed.acceptance_start.isoformat() == "2026-04-09"
    assert parsed.acceptance_end.isoformat() == "2026-04-10"
    assert parsed.has_invalid_acceptance_range is False


def test_parse_docx_marks_reversed_acceptance_range_invalid():
    text = "竣工验收2026年04月10日2026年04月09日"

    parsed = parse_docx_text(text)

    assert parsed.has_invalid_acceptance_range is True
    assert parsed.acceptance_start.isoformat() == "2026-04-10"
    assert parsed.acceptance_end.isoformat() == "2026-04-09"


def test_parse_docx_collects_multiple_owner_contacts():
    text = (
        "项目编号：BHE-25030367/01 "
        "用户姓名 张三 "
        "联系电话 13800000000 "
        "用户姓名 黄汉民 "
        "联系电话 14714691425 "
        "竣工验收2026年04月09日2026年04月10日"
    )

    parsed = parse_docx_text(text)

    assert parsed.contact_names == ["张三", "黄汉民"]
    assert parsed.contact_phones == ["13800000000", "14714691425"]


def test_read_docx_text_includes_table_cells(tmp_path: Path):
    path = tmp_path / "table.docx"
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目全称"
    table.cell(0, 1).text = "示例项目"
    table.cell(1, 0).text = "用户姓名"
    table.cell(1, 1).text = "黄汉民"
    document.save(path)

    text = read_docx_text(path)

    assert "项目全称" in text
    assert "示例项目" in text
    assert "用户姓名" in text
    assert "黄汉民" in text


def test_read_doc_text_supports_plain_text_doc(tmp_path: Path):
    path = tmp_path / "项目竣工总结报告.doc"
    path.write_text(
        "项目编号：BHE-25030367/01 项目全称 示例项目 用户姓名 黄汉民 联系电话14714691425",
        encoding="utf-8",
    )

    text = read_doc_text(path)

    assert "项目编号" in text
    assert "BHE-25030367/01" in text
